# SPDX-FileCopyrightText: 2026 Department of Health – Western Visayas Center for Health Development
# SPDX-License-Identifier: GPL-3.0-or-later
"""Build the web scorecard: copy tool/ and write the criteria data it reads.

    python scripts/build_web.py [--out site]

Open site/index.html through a local web server to try it, e.g.:
    python -m http.server -d site 8000
"""
from __future__ import annotations

import argparse
import json
import os
import shutil
from pathlib import Path

from build_reference import build_reference
from build_xlsx import retired_list
from common import CODE_LICENSE, CONTENT_LICENSE, CREDIT, OWNER, ROOT, TEMPLATE_XLSX, TITLE, active, active_subs, load_criteria, load_scoring, load_sheets, version_info


def report_text() -> dict:
    """Background and scope paragraphs of the printed report, kept in the workbook template."""
    import warnings
    import openpyxl
    warnings.filterwarnings("ignore")
    ws = openpyxl.load_workbook(TEMPLATE_XLSX)["Printed Evaluation Report"]
    out = {}
    for r in range(1, 40):
        v = ws.cell(r, 1).value
        if v in ("Background", "Scope and Coverage"):
            out["background" if v == "Background" else "scope"] = str(ws.cell(r + 1, 1).value or "").replace("\r\n", "\n")
    return out


def export(out: Path) -> Path:
    info = version_info()
    sheets = load_sheets()
    criteria = load_criteria()
    scoring = load_scoring()
    repo = os.environ.get("GITHUB_REPOSITORY", "")
    data = {
        "meta": {**info, "repository": f"https://github.com/{repo}" if repo else "", "owner": OWNER,
                 "content_license": CONTENT_LICENSE, "code_license": CODE_LICENSE, "credit": CREDIT, "title": TITLE},
        "report": report_text(),
        "scoring": scoring,
        "sheets": [
            {
                "key": s["key"], "tab": s["tab"], "title": s["title"], "intro": s["intro"],
                "threshold": scoring["thresholds"][s["key"]],
                "criteria": [
                    {k: v for k, v in c.items() if not k.startswith("_") and k not in ("status",)}
                    | {"sub_items": active_subs(c)}
                    for c in active(criteria[s["key"]])
                ],
            }
            for s in sheets
        ],
        "retired": retired_list(criteria),
    }
    if out.exists():
        shutil.rmtree(out)
    shutil.copytree(ROOT / "tool", out, ignore=shutil.ignore_patterns("landing.html"))
    (out / "data.json").write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")
    (out / "reference.json").write_text(json.dumps(build_reference(), ensure_ascii=False), encoding="utf-8")
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", type=Path, default=ROOT / "site")
    args = ap.parse_args()
    print(f"Wrote web scorecard to {export(args.out)}")


if __name__ == "__main__":
    main()
