"""Generate the scorecard workbook from the criteria files.

    python scripts/build_xlsx.py [--out dist]

Uses templates/scorecard-template.xlsx for layout, styles and the sheets that are not
generated (Instructions, System Profile, Test Data). Regenerates: every criteria sheet,
the Summary thresholds, policy and retired-ID table, and the criteria breakdown of the
Printed Evaluation Report. Stamps version and commit on every sheet.
"""
from __future__ import annotations

import argparse
import copy
import math
import re
import warnings
from pathlib import Path

import openpyxl
from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.worksheet.datavalidation import DataValidation

from common import (CRITICAL_TO_XLSX, ROOT, SCORING_TO_XLSX, TEMPLATE_XLSX, active, active_subs, load_criteria,
                    load_legacy_retired, load_scoring, load_sheets, version_info)

warnings.filterwarnings("ignore")

FIRST = 10            # first criteria row on a criteria sheet
NCOLS = 23            # A..W
REPORT_SHEET = "Printed Evaluation Report"
SUMMARY_ROWS = {"L0": 7, "L1": 8, "L2A": 10, "L2B": 11, "L2C": 12, "L2D": 13, "L3": 14, "L4": 15}


# ----------------------------------------------------------------------------- formulas
def f_score(r):
    return (f'=IF(M{r}="N/A",IF(H{r}="Yes",0,""),IF(M{r}="Not Tested",0,IF(J{r}="Tier",IF(M{r}="Fully integrated",1,'
            f'IF(M{r}="Standalone - manual encoding",0.5,0)),IF(J{r}="Binary",IF(M{r}="Met",1,0),IF(M{r}="Counted",'
            f'IF(N(K{r})>0,IF(J{r}="All-or-nothing",IF(N(L{r})>=K{r},1,0),MIN(N(L{r})/K{r},1)),0),0)))))')


def f_outstanding_simple(r):
    return f'=IF(M{r}="Not Met","Needs compliance",IF(M{r}="Not Tested","Not yet tested",""))'


def f_needs(r):
    return f'=IF(AND(T{r}="Criterion",I{r}>0,M{r}<>"Not Tested",N{r}<>"",OR(N(N{r})<1,S{r}=1)),A{r},"")'


def criterion_row(c, r, a=None, b=None):
    """Values and formulas for a criterion row r; a..b = its sub-item rows."""
    row = {
        "A": c["id"], "B": c["section"], "C": c["title"], "D": c["developer_guide"], "E": c["validator_guide"],
        "F": c["evidence"], "G": c["validation_role"], "H": CRITICAL_TO_XLSX[c["critical"]], "I": c["weight"],
        "J": SCORING_TO_XLSX[c["scoring"]], "N": f_score(r), "Q": f'=IF(N{r}="",0,I{r})',
        "R": f'=IF(N{r}="",0,I{r}*N{r})', "T": "Criterion", "W": f_needs(r),
    }
    if a is None:
        row["M"] = "Not Tested"
        row["O"] = f_outstanding_simple(r) if c["scoring"] != "tier" else (
            f'=IF(M{r}="Not Met","Needs compliance",IF(M{r}="Not Tested","Not yet tested",'
            f'IF(N{r}<1,"Integration level: "&M{r},"")))')
        row["S"] = f'=IF(AND(H{r}="Yes",M{r}<>"Not Tested",N{r}<>"",N{r}<1),1,0)'
    else:
        rng = f"{a}:A{b}"
        row["K"] = f'=COUNTA(A{a}:A{b})-COUNTIF(M{a}:M{b},"N/A")'
        row["L"] = f'=COUNTIF(M{a}:M{b},"Met")'
        row["M"] = (f'=IF(COUNTIF(M{a}:M{b},"N/A")=COUNTA(A{a}:A{b}),"N/A",IF(COUNTIF(M{a}:M{b},"Not Tested")>0,'
                    f'"Not Tested","Counted"))')
        row["O"] = (f'=IF(COUNTIF(M{a}:M{b},"Not Met")>0,COUNTIF(M{a}:M{b},"Not Met")&" of "&K{r}&" need compliance: "'
                    f'&_xlfn.TEXTJOIN("; ",TRUE(),U{a}:U{b}),IF(COUNTIF(M{a}:M{b},"Not Tested")>0,'
                    f'COUNTIF(M{a}:M{b},"Not Tested")&" of "&COUNTA(A{a}:A{b})&" not yet tested",'
                    f'IF(M{r}="N/A","","All compliant")))')
        row["S"] = (f'=IF(OR(AND(H{r}="Yes",M{r}<>"Not Tested",N{r}<>"",N{r}<1),COUNTIFS(H{a}:H{b},"Yes",M{a}:M{b},'
                    f'"Not Met")+COUNTIFS(H{a}:H{b},"Yes",M{a}:M{b},"N/A")>0),1,0)')
        row["V"] = f'=COUNTA(A{a}:A{b})-COUNTIF(M{a}:M{b},"Not Tested")'
        del rng
    return row


def sub_row(s, parent_id, r, n):
    row = {
        "A": s["id"], "B": "↳ sub-item", "C": s["text"], "D": f"(part of {parent_id})", "M": "Not Tested",
        "O": f_outstanding_simple(r), "T": "Sub",
        "U": f'=IF(M{r}="Not Met","("&{n}&") "&IF(H{r}="Yes","★ ","")&C{r},"")',
    }
    if s["critical"]:
        row["H"] = "Yes"
    return row


# ----------------------------------------------------------------------------- helpers
def row_styles(ws, r):
    return {c: copy.copy(ws.cell(r, c)._style) for c in range(1, NCOLS + 1)}


def apply_styles(ws, r, styles, ncols=NCOLS):
    for c in range(1, ncols + 1):
        ws.cell(r, c)._style = copy.copy(styles[c])


def find_rows(ws, pred, start=FIRST):
    return [r for r in range(start, ws.max_row + 1) if pred(r)]


def unmerge_from(ws, first_row):
    for m in list(ws.merged_cells.ranges):
        if m.min_row >= first_row:
            ws.unmerge_cells(str(m))


def merges_of_row(ws, r):
    return [(m.min_col, m.max_col) for m in ws.merged_cells.ranges if m.min_row == r and m.max_row == r]


def lines_needed(text, width_chars):
    if not text:
        return 1
    total = 0
    for part in str(text).split("\n"):
        total += max(1, math.ceil(len(part) / max(width_chars, 1)))
    return total


def col_width(ws, letter, default=10):
    w = ws.column_dimensions[letter].width
    return w if w else default


def resize_cf(ws, first_row, last_row):
    """Re-anchor conditional formats that cover the generated rows."""
    old = ws.conditional_formatting
    new = ConditionalFormattingList()
    for cf in old:
        ranges = []
        for rng in str(cf.sqref).split():
            m = re.match(r"([A-Z]+)(\d+)(?::([A-Z]+)(\d+))?$", rng)
            if m and int(m.group(2)) >= first_row:
                ranges.append(f"{m.group(1)}{m.group(2)}:{m.group(3) or m.group(1)}{last_row}")
            else:
                ranges.append(rng)
        for rule in cf.rules:
            new.add(" ".join(ranges), rule)
    ws.conditional_formatting = new


# ----------------------------------------------------------------------------- criteria sheets
def build_criteria_sheet(ws, sheet, criteria):
    # capture styles from the template rows before clearing
    t_crit = next(r for r in find_rows(ws, lambda r: ws.cell(r, 20).value == "Criterion"))
    t_sub = next(r for r in find_rows(ws, lambda r: ws.cell(r, 20).value == "Sub"))
    crit_style, sub_style = row_styles(ws, t_crit), row_styles(ws, t_sub)
    dv_template = ws.data_validations.dataValidation[0] if ws.data_validations.dataValidation else None

    unmerge_from(ws, FIRST)
    ws.delete_rows(FIRST, ws.max_row)
    for r in [r for r in list(ws.row_dimensions) if r >= FIRST]:
        del ws.row_dimensions[r]

    ws["A1"] = sheet["title"]
    ws["A2"] = sheet["intro"]
    if sheet.get("note"):
        ws["G4"] = sheet["note"]

    widths = {L: col_width(ws, L) for L in "ABCDEFGHIJKLMNOP"}
    standard_cells, tier_cells = [], []
    r = FIRST
    for c in active(criteria):
        subs = active_subs(c)
        a, b = (r + 1, r + len(subs)) if subs else (None, None)
        values = criterion_row(c, r, a, b)
        apply_styles(ws, r, crit_style)
        for col, v in values.items():
            ws[f"{col}{r}"] = v
        if not subs:
            (tier_cells if c["scoring"] == "tier" else standard_cells).append(f"M{r}")
        height = max(lines_needed(values[k], widths[k] * 1.15) for k in "BCDEF") * 13.5 + 4
        ws.row_dimensions[r].height = min(max(height, 15), 409)
        r += 1
        for n, s in enumerate(subs, 1):
            values = sub_row(s, c["id"], r, n)
            apply_styles(ws, r, sub_style)
            for col, v in values.items():
                ws[f"{col}{r}"] = v
            standard_cells.append(f"M{r}")
            ws.row_dimensions[r].height = max(lines_needed(s["text"], widths["C"] * 1.15) * 13.5 + 2, 15)
            r += 1
    last = r - 1

    # data validation for the Result column
    ws.data_validations.dataValidation = []
    for cells, choices in ((standard_cells, '"Met,Not Met,Not Tested,N/A"'),
                           (tier_cells, '"Fully integrated,Standalone - manual encoding,Not Met,Not Tested"')):
        if not cells:
            continue
        dv = DataValidation(type="list", formula1=choices, allow_blank=True)
        if dv_template is not None:
            for attr in ("showErrorMessage", "showInputMessage", "errorTitle", "error", "promptTitle", "prompt", "errorStyle"):
                setattr(dv, attr, getattr(dv_template, attr))
        dv.sqref = openpyxl.worksheet.cell_range.MultiCellRange(" ".join(cells))
        ws.add_data_validation(dv)

    resize_cf(ws, FIRST, last)
    ws.auto_filter.ref = f"A9:P{last}"
    ws.print_area = f"A1:P{last}"
    return last


# ----------------------------------------------------------------------------- summary
def build_summary(ws, scoring, retired_rows):
    for key, row in SUMMARY_ROWS.items():
        ws.cell(row, 4).value = scoring["thresholds"][key]
    pol = scoring["policy"]
    ws["C23"], ws["C24"], ws["C25"], ws["C26"] = pol["status"], pol["approved_by"] or None, pol["date_approved"] or None, pol["reference"] or None

    # retired table starts at row 43; the "HOW TO READ" block follows it
    data_style = {c: copy.copy(ws.cell(43, c)._style) for c in range(1, 10)}
    data_merges = merges_of_row(ws, 43)
    how_start = next(r for r in range(43, ws.max_row + 1) if ws.cell(r, 2).value == "HOW TO READ THIS SHEET")
    how = [([ws.cell(r, c).value for c in range(1, 10)], {c: copy.copy(ws.cell(r, c)._style) for c in range(1, 10)},
            merges_of_row(ws, r), ws.row_dimensions[r].height) for r in range(how_start, ws.max_row + 1)]
    unmerge_from(ws, 43)
    for r in range(43, ws.max_row + 1):
        for c in range(1, 10):
            ws.cell(r, c).value = None
        if r in ws.row_dimensions:
            ws.row_dimensions[r].height = None

    r = 43
    for item in retired_rows:
        for c in range(1, 10):
            ws.cell(r, c)._style = copy.copy(data_style[c])
        ws.cell(r, 2).value, ws.cell(r, 3).value, ws.cell(r, 4).value = item["id"], item["revision"], item["action"]
        ws.cell(r, 5).value, ws.cell(r, 6).value = item["now"], item["reason"]
        for a, b in data_merges:
            ws.merge_cells(start_row=r, start_column=a, end_row=r, end_column=b)
        r += 1
    r += 1
    for values, styles, merges, height in how:
        for c in range(1, 10):
            ws.cell(r, c)._style = copy.copy(styles[c])
            ws.cell(r, c).value = values[c - 1]
        for a, b in merges:
            ws.merge_cells(start_row=r, start_column=a, end_row=r, end_column=b)
        ws.row_dimensions[r].height = height
        r += 1


def retired_list(criteria_by_sheet):
    rows = [{"id": x["id"], "revision": x["legacy_revision"], "action": x["action"],
             "now": x.get("replaced_by") or "—", "reason": x["reason"]} for x in load_legacy_retired()]
    for items in criteria_by_sheet.values():
        for c in items:
            if c.get("status", "active") != "active":
                rows.append({"id": c["id"], "revision": f"v{c.get('retired_in', '')}", "action": c["status"].capitalize(),
                             "now": c.get("replaced_by") or "—", "reason": c.get("reason", "")})
            for s in c.get("sub_items") or []:
                if s.get("status", "active") != "active":
                    rows.append({"id": s["id"], "revision": f"v{s.get('retired_in', '')}", "action": "Removed",
                                 "now": "—", "reason": s.get("reason", "")})
    return rows


# ----------------------------------------------------------------------------- printed report
def build_report(ws, sheets, criteria_by_sheet):
    start = next(r for r in range(1, ws.max_row + 1) if str(ws.cell(r, 1).value or "").startswith("LEVEL 0"))
    t_head, t_score = start, start + 1
    t_crit = next(r for r in range(start, ws.max_row + 1) if str(ws.cell(r, 2).value or "").startswith("='L0 Application'!C"))
    t_guide = t_crit + 1
    t_sub = next(r for r in range(start, ws.max_row + 1) if str(ws.cell(r, 2).value or "").startswith('="   ↳'))

    def capture(r):
        return ({c: copy.copy(ws.cell(r, c)._style) for c in range(1, 7)},
                [ws.cell(r, c).value for c in range(1, 7)], merges_of_row(ws, r), ws.row_dimensions[r].height)

    head, score, crit, guide, sub = (capture(x) for x in (t_head, t_score, t_crit, t_guide, t_sub))
    src_sheet, src_crit_row, src_sub_row = "'L0 Application'!", 10, 12
    widths = [col_width(ws, L, 12) for L in "ABCDEF"]

    def localise(formula, tab, src_row=None, row=None):
        if not isinstance(formula, str):
            return formula
        out = formula.replace(src_sheet, f"'{tab}'!")
        if src_row is not None:
            out = re.sub(rf"('{re.escape(tab)}'!\$?[A-Z]{{1,2}}\$?){src_row}\b", rf"\g<1>{row}", out)
        return out

    unmerge_from(ws, start)
    for r in range(start, ws.max_row + 1):
        for c in range(1, 7):
            ws.cell(r, c).value = None
    for r in [r for r in list(ws.row_dimensions) if r >= start]:
        del ws.row_dimensions[r]

    def put(r, block, values, height=None):
        styles, _, merges, h = block
        for c in range(1, 7):
            ws.cell(r, c)._style = copy.copy(styles[c])
            ws.cell(r, c).value = values[c - 1]
        for a, b in merges:
            ws.merge_cells(start_row=r, start_column=a, end_row=r, end_column=b)
        ws.row_dimensions[r].height = height if height is not None else h

    def span_width(block, col):
        for a, b in block[2]:
            if a <= col <= b:
                return sum(widths[a - 1:b])
        return widths[col - 1]

    r = start
    for sheet in sheets:
        tab = sheet["tab"]
        put(r, head, [sheet.get("report_heading") or sheet["title"].upper(), None, None, None, None, None])
        r += 1
        put(r, score, [localise(v, tab) for v in score[1]])
        r += 1
        src_row = FIRST
        for c in active(criteria_by_sheet[sheet["key"]]):
            # the criterion's row on its own sheet
            text = c["title"] + "\n[" + c["section"] + "]"
            h = max(lines_needed(text, span_width(crit, 2) * 1.1), 2) * 13 + 4
            put(r, crit, [localise(v, tab, src_crit_row, src_row) for v in crit[1]], h)
            r += 1
            dev = c["developer_guide"].split("\n\nPurpose:")[0]
            gh = max(lines_needed("Developer's Guide: " + dev, span_width(guide, 2) * 1.1),
                     lines_needed("Validator's Guide: " + c["validator_guide"] + "\nEvidence: " + c["evidence"],
                                  span_width(guide, 3) * 1.1)) * 12 + 4
            put(r, guide, [localise(v, tab, src_crit_row, src_row) for v in guide[1]], min(gh, 409))
            r += 1
            src_row += 1
            for s in active_subs(c):
                sh = lines_needed("   ↳  ★ " + s["text"], span_width(sub, 2) * 1.1) * 13 + 3
                put(r, sub, [localise(v, tab, src_sub_row, src_row) for v in sub[1]], sh)
                r += 1
                src_row += 1
    last = r - 1
    resize_cf(ws, start, last)
    ws.print_area = f"A1:F{last}"


# ----------------------------------------------------------------------------- main
def stamp(wb, info, policy_status):
    label = f"Version {info['label']} · built {info['built']}"
    for ws in wb.worksheets:
        ws.oddFooter.left.text = label
        ws.oddFooter.left.size = 8
    ins = wb["Instructions"]
    cur = ins["B3"].value or ""
    rest = re.sub(r"^DRAFT — Revision \d+ \([^)]*\)\.\s*", "", cur)
    rest = re.sub(r"^Version [^.]*\.\s*", "", rest)
    prefix = "" if info["is_release"] else "DRAFT — "
    ins["B3"] = f"{prefix}Version {info['label']} (built {info['built']}). " + rest.replace(
        "Thresholds and policy values are PROPOSED", f"Thresholds and policy values are {policy_status.upper()}")


def build(out_dir: Path) -> Path:
    info = version_info()
    sheets = load_sheets()
    scoring = load_scoring()
    criteria = load_criteria()
    wb = openpyxl.load_workbook(TEMPLATE_XLSX)
    for sheet in sheets:
        build_criteria_sheet(wb[sheet["tab"]], sheet, criteria[sheet["key"]])
    build_summary(wb["Summary"], scoring, retired_list(criteria))
    build_report(wb[REPORT_SHEET], sheets, criteria)
    stamp(wb, info, scoring["policy"]["status"])
    out_dir.mkdir(parents=True, exist_ok=True)
    safe = re.sub(r"[^A-Za-z0-9._-]+", "-", info["version"]).strip("-")
    path = out_dir / f"testing-criteria-scorecard-{safe}.xlsx"
    wb.save(path)
    return path


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", type=Path, default=ROOT / "dist")
    args = ap.parse_args()
    print(f"Wrote {build(args.out)}")


if __name__ == "__main__":
    main()
