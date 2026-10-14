"""Build the workbook and check every criterion row matches the YAML files exactly.

    python tests/test_roundtrip.py

Fast (no LibreOffice needed); runs on every push.
"""
from __future__ import annotations

import sys
import tempfile
import warnings
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))
from build_xlsx import build  # noqa: E402
from common import CRITICAL_TO_XLSX, SCORING_TO_XLSX, active, active_subs, load_criteria, load_sheets  # noqa: E402

warnings.filterwarnings("ignore")


def main():
    errors = []
    criteria = load_criteria()
    with tempfile.TemporaryDirectory() as t:
        wb = openpyxl.load_workbook(build(Path(t)))
        for s in load_sheets():
            ws = wb[s["tab"]]
            expected = []
            for c in active(criteria[s["key"]]):
                expected.append(("Criterion", c["id"], {
                    "B": c["section"], "C": c["title"], "D": c["developer_guide"], "E": c["validator_guide"], "F": c["evidence"],
                    "G": c["validation_role"], "H": CRITICAL_TO_XLSX[c["critical"]], "I": c["weight"], "J": SCORING_TO_XLSX[c["scoring"]]}))
                for x in active_subs(c):
                    expected.append(("Sub", x["id"], {"C": x["text"], "H": "Yes" if x["critical"] else None}))
            rows = [r for r in range(10, ws.max_row + 1) if ws.cell(r, 1).value]
            if len(rows) != len(expected):
                errors.append(f"{s['tab']}: {len(rows)} rows in workbook, {len(expected)} expected")
                continue
            for r, (kind, i, fields) in zip(rows, expected):
                if ws.cell(r, 1).value != i or ws.cell(r, 20).value != kind:
                    errors.append(f"{s['tab']} row {r}: expected {kind} {i}, found {ws.cell(r, 20).value} {ws.cell(r, 1).value}")
                    continue
                for col, v in fields.items():
                    if ws[f"{col}{r}"].value != v:
                        errors.append(f"{s['tab']} {col}{r} ({i}): workbook {ws[f'{col}{r}'].value!r} != YAML {v!r}")
    for e in errors[:50]:
        print("ERROR", e)
    print(f"Round trip: {'ok' if not errors else str(len(errors)) + ' difference(s)'}")
    return 1 if errors else 0


if __name__ == "__main__":
    sys.exit(main())
