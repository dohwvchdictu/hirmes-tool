"""Check that the workbook and the web scorecard score the same answers identically.

    python tests/test_scoring.py

For each scenario in tests/scenarios.yaml the script fills the answers into a freshly built
workbook, recalculates it with LibreOffice (headless), scores the same answers with
tool/scoring.js under Node.js, and compares every sheet and level result.
Needs: LibreOffice (soffice) and Node.js on the PATH.
"""
from __future__ import annotations

import json
import random
import shutil
import subprocess
import sys
import tempfile
import warnings
from pathlib import Path

import openpyxl
import yaml

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))
from build_web import export  # noqa: E402
from build_xlsx import build  # noqa: E402
from common import active, active_subs, load_criteria, load_scoring, load_sheets  # noqa: E402

warnings.filterwarnings("ignore")
CHOICES = ["Met", "Not Met", "Not Tested", "N/A"]
TIER = ["Fully integrated", "Standalone - manual encoding", "Not Met", "Not Tested"]


def answer_slots():
    """[(sheet_key, tab, id, is_tier)] in workbook order."""
    out = []
    for s in load_sheets():
        for c in active(load_criteria()[s["key"]]):
            subs = active_subs(c)
            if subs:
                out += [(s["key"], s["tab"], x["id"], False) for x in subs]
            else:
                out.append((s["key"], s["tab"], c["id"], c["scoring"] == "tier"))
    return out


def make_answers(sc, slots):
    rnd = random.Random(sc.get("seed", 0))
    weights = sc.get("weights", [1, 1, 1, 1])
    ans = {}
    for _, _, i, tier in slots:
        if "all" in sc:
            v = sc["all"]
            if tier:
                v = {"Met": "Fully integrated", "N/A": "Not Met"}.get(v, v)
        else:
            v = rnd.choices(TIER if tier else CHOICES, weights=weights)[0]
        if v != "Not Tested":
            ans[i] = v
    return ans


def fill_and_recalc(xlsx: Path, answers: dict, tmp: Path) -> Path:
    wb = openpyxl.load_workbook(xlsx)
    for s in load_sheets():
        ws = wb[s["tab"]]
        for r in range(10, ws.max_row + 1):
            i = ws.cell(r, 1).value
            if i in answers and isinstance(ws.cell(r, 13).value, str) and not ws.cell(r, 13).value.startswith("="):
                ws.cell(r, 13).value = answers[i]
    tmp.mkdir(parents=True, exist_ok=True)
    filled = tmp / "filled.xlsx"
    wb.save(filled)
    outdir = tmp / "calc"
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", str(outdir), str(filled)],
                   check=True, capture_output=True, env={"HOME": str(tmp), "PATH": "/usr/bin:/bin:/usr/local/bin"})
    return outdir / "filled.xlsx"


def read_xlsx(path: Path):
    wb = openpyxl.load_workbook(path, data_only=True)
    sheets = {}
    for s in load_sheets():
        ws = wb[s["tab"]]
        sheets[s["key"]] = {"score": round(ws["D4"].value or 0, 6), "criticalNotMet": ws["D6"].value, "notTested": ws["F4"].value,
                            "status": ws["F6"].value, "total": ws["F7"].value}
    sm = wb["Summary"]
    gated = {k: sm.cell(r, 9).value for k, r in (("L0", 7), ("L1", 8), ("L2", 9), ("L3", 14), ("L4", 15))}
    return {"sheets": sheets, "gated": gated, "highest": sm["C18"].value}


def read_js(site: Path, answers: dict):
    js = (
        "const S=require(process.argv[1]);const d=require(process.argv[2]);const a=JSON.parse(process.argv[3]);"
        "const r=S.scoreAll(d,a);const o={sheets:{},gated:{},highest:r.highest};"
        "for(const k in r.sheets){const x=r.sheets[k];o.sheets[k]={score:Math.round(x.score*1e6)/1e6,criticalNotMet:x.criticalNotMet,"
        "notTested:x.notTested,status:x.status,total:x.total};}"
        "for(const l of r.levels){o.gated[l.key]=l.gated;}console.log(JSON.stringify(o));"
    )
    out = subprocess.run(["node", "-e", js, str(site / "scoring.js"), str(site / "data.json"), json.dumps(answers)],
                         capture_output=True, text=True, check=True).stdout
    return json.loads(out)


def main():
    if not shutil.which("soffice") or not shutil.which("node"):
        print("SKIPPED: needs LibreOffice (soffice) and Node.js")
        return 0
    scenarios = yaml.safe_load((ROOT / "tests" / "scenarios.yaml").read_text())["scenarios"]
    slots = answer_slots()
    failures = 0
    with tempfile.TemporaryDirectory() as t:
        tmp = Path(t)
        xlsx = build(tmp / "dist")
        site = export(tmp / "site")
        for sc in scenarios:
            answers = make_answers(sc, slots)
            x = read_xlsx(fill_and_recalc(xlsx, answers, tmp / sc["name"]))
            j = read_js(site, answers)
            diffs = []
            for k, v in x["sheets"].items():
                for f, xv in v.items():
                    if xv != j["sheets"][k][f]:
                        diffs.append(f"{k}.{f}: workbook {xv} vs web {j['sheets'][k][f]}")
            for k, v in x["gated"].items():
                if v != j["gated"][k]:
                    diffs.append(f"level {k}: workbook {v} vs web {j['gated'][k]}")
            if x["highest"] != j["highest"]:
                diffs.append(f"highest: workbook {x['highest']} vs web {j['highest']}")
            status = "ok" if not diffs else "MISMATCH"
            print(f"{sc['name']:<22} {status}  highest={x['highest']}  " + " ".join(f"{k}:{v['status']}" for k, v in x["sheets"].items()))
            for d in diffs:
                print("    " + d)
            failures += bool(diffs)
    print(f"\n{len(scenarios) - failures} of {len(scenarios)} scenarios agree.")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
